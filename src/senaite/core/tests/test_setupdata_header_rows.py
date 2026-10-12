# -*- coding: utf-8 -*-
import os
import unittest

from openpyxl import Workbook
from openpyxl import load_workbook
from senaite.core.exportimport.setupdata import WorksheetImporter


class TestSetupDataHeaderRows(unittest.TestCase):

    def rows(self, name, rows, startrow=3):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = name
        for row in rows:
            worksheet.append(row)
        importer = WorksheetImporter(None)
        importer.worksheet = worksheet
        return list(importer.get_rows(startrow))

    def test_sample_points_two_headers_keep_first_record(self):
        rows = self.rows("Sample Points", [
            ["Client_title", "title", "SampleType_title"],
            ["Client", "Sample Point", "Sample Type"],
            ["Plant", "KUM-FDP-MFP", "Solid (crushed ore)"],
            ["Plant", "KUM-MIL-CFD", "Slurry"],
        ])
        self.assertEqual([row["title"] for row in rows],
                         ["KUM-FDP-MFP", "KUM-MIL-CFD"])

    def test_sample_types_two_headers_keep_resin(self):
        rows = self.rows("Sample Types", [
            ["title", "Prefix"], ["Sample Type", "Prefix"],
            ["Resin", "RSN"], ["Slurry", "SLR"],
        ])
        self.assertEqual([row["title"] for row in rows], ["Resin", "Slurry"])

    def test_legacy_three_headers_are_preserved(self):
        for name, label in (("Sample Points", "Sample Point"),
                            ("Sample Types", "Sample Type")):
            rows = self.rows(name, [
                ["title"], ["Description of sheet"], [label], ["First record"],
            ])
            self.assertEqual([row["title"] for row in rows], ["First record"])

    def test_compact_one_header_layout_is_preserved(self):
        rows = self.rows("Sample Points", [
            ["Client", "Sample Point", "Sample Type"],
            ["Plant", "KUM-FDP-MFP", "Solid"],
        ], startrow=1)
        self.assertEqual(rows[0]["Sample Point"], "KUM-FDP-MFP")

    def test_lotus_v2_first_records_are_preserved(self):
        path = os.path.abspath(os.path.join(
            os.path.dirname(__file__), "../../../../../bika.mining/src/bika/mining",
            "setupdata/Lotus new setup data v2.xlsx"))
        if not os.path.isfile(path):
            self.skipTest("Mining workbook is not available in this checkout")
        workbook = load_workbook(path)
        for name, title in (("Sample Points", "KUM-FDP-MFP"),
                            ("Sample Types", "Resin")):
            importer = WorksheetImporter(None)
            importer.worksheet = workbook[name]
            rows = list(importer.get_rows(3))
            self.assertEqual(rows[0]["title"], title)


def test_suite():
    return unittest.TestLoader().loadTestsFromTestCase(TestSetupDataHeaderRows)
